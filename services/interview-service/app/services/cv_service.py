"""Reads an uploaded CV and turns it into an interview profile.

The CV is parsed in memory and never stored. The profile (skills, suggested role, experience
level, short background) seeds the interview: the role and level preselect the dropdowns and the
background personalises the questions. Uses the configured LLM; without one it falls back to
keyword matching and says so in `method`.
"""
import io
import re
from typing import Any, Dict, List

from app.core.config import settings

MAX_CV_BYTES = 5 * 1024 * 1024
MAX_CV_CHARS = 12000                       # enough for a long CV; keeps the local LLM prompt small
LEVELS = ["Entry Level", "Mid Level", "Senior Level"]

# Keyword fallback: role -> terms that suggest it
ROLE_KEYWORDS = {
    "Backend Developer": ["java", "spring", "node", "express", "django", "flask", "fastapi", "rest", "api",
                          "microservice", "kafka", "rabbitmq", "redis", "go", "golang", "c#", ".net", "php"],
    "Full-Stack Developer": ["react", "next.js", "nextjs", "vue", "angular", "typescript", "javascript",
                             "html", "css", "tailwind", "frontend", "full-stack", "full stack", "mern"],
    "DevOps Engineer": ["docker", "kubernetes", "terraform", "ansible", "jenkins", "github actions", "ci/cd",
                        "aws", "azure", "gcp", "linux", "prometheus", "grafana", "helm", "devops"],
    "Database Engineer": ["sql", "postgresql", "postgres", "mysql", "oracle", "mongodb", "database",
                          "indexing", "query optimization", "data modeling", "etl", "dba"],
}
# Non-technical roles bring their own keywords (app/core/role_profiles.json)
from app.core.roles import non_technical_profiles  # noqa: E402
# Keywords are matched as substrings, so "hr " and "ca " keep their trailing space
ROLE_KEYWORDS.update({role: p["cv_keywords"] for role, p in non_technical_profiles().items()})
SKILL_VOCAB = sorted({k for kws in ROLE_KEYWORDS.values() for k in kws} |
                     {"python", "c++", "git", "graphql", "jwt", "oauth", "spring boot", "hibernate",
                      "pandas", "machine learning", "flutter", "android", "kotlin", "swift"})


class CVError(Exception):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_CV_BYTES:
        raise CVError("The CV is larger than 5 MB.", status_code=413)
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or data[:4] == b"%PDF":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
        elif name.endswith(".docx"):
            import docx
            document = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="ignore")
        else:
            raise CVError("Upload a PDF, DOCX or TXT file.", status_code=415)
    except CVError:
        raise
    except Exception as e:
        raise CVError(f"Couldn't read the CV ({type(e).__name__}). Try another format.") from e
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 40:
        raise CVError("No readable text was found in the CV. If it's a scanned image, upload a text-based PDF or DOCX.")
    return text[:MAX_CV_CHARS]


def _stated_level(text: str) -> str | None:
    """Experience level from facts the CV states outright (years of experience, student status).
    These override the LLM's guess, which in testing called a 6-year DBA and a student 'Mid Level'."""
    low = text.lower()
    years = [int(y) for y in re.findall(r"(\d{1,2})\+?\s*(?:years|yrs)", low)]
    if years and max(years) >= 5:
        return "Senior Level"
    if years and max(years) >= 2:
        return "Mid Level"
    if re.search(r"\b(undergraduate|student|final[- ]year|fresh graduate|graduate|intern)\b", low):
        return "Entry Level"
    return None


def _keyword_profile(text: str, roles: List[str]) -> Dict[str, Any]:
    low = text.lower()
    skills = [s.strip() for s in SKILL_VOCAB if re.search(r"(?<![a-z])" + re.escape(s.strip()) + r"(?![a-z])", low)]
    scores = {r: sum(1 for k in ROLE_KEYWORDS.get(r, []) if k in low) for r in roles}
    role = max(scores, key=scores.get) if scores and max(scores.values()) > 0 else (roles[0] if roles else "")
    level = _stated_level(text) or "Entry Level"
    return {"skills": skills[:15], "suggested_role": role, "experience_level": level,
            "summary": "", "method": "keyword"}


PROFILE_PROMPT = (
    "You read a candidate's CV for a mock job interview. Return a JSON object with keys:\n"
    "- 'name': the candidate's first name, or '' if unclear\n"
    "- 'skills': list of up to 15 concrete skills: technical (languages, frameworks, tools) or professional "
    "(e.g. recruitment, SEO, budgeting, copywriting, Adobe Illustrator, brand strategy)\n"
    "- 'suggested_role': exactly one of: {roles}\n"
    "- 'experience_level': exactly one of: Entry Level, Mid Level, Senior Level "
    "(students and graduates are Entry Level)\n"
    "- 'summary': one sentence describing their background and strongest areas, used to tailor questions\n"
    "Use only what the CV says; do not invent experience."
)


async def analyse_cv(text: str, roles: List[str]) -> Dict[str, Any]:
    roles = roles or list(ROLE_KEYWORDS)
    system = PROFILE_PROMPT.format(roles=", ".join(roles))
    data = None
    try:
        if settings.LLM_PROVIDER == "local":
            from app.services.local_llm import local_llm
            data = await local_llm.chat_json(system, f"CV:\n{text}", temperature=0.1, max_tokens=500)
        elif settings.LLM_PROVIDER == "gemini" and settings.GEMINI_API_KEY:
            import json
            from google import genai
            from google.genai import types
            client = genai.Client(api_key=settings.GEMINI_API_KEY)
            resp = client.models.generate_content(
                model=settings.LLM_MODEL or "gemini-2.5-flash", contents=f"{system}\n\nCV:\n{text}",
                config=types.GenerateContentConfig(temperature=0.1, response_mime_type="application/json"))
            data = json.loads(resp.text)
    except Exception as e:
        print(f"[Warning] CV analysis with LLM failed, using keyword matching: {e}")

    fallback = _keyword_profile(text, roles)
    if not isinstance(data, dict):
        return fallback
    skills = [str(s).strip() for s in (data.get("skills") or []) if str(s).strip()][:15]
    role = data.get("suggested_role") if data.get("suggested_role") in roles else fallback["suggested_role"]
    level = _stated_level(text) or (data.get("experience_level") if data.get("experience_level") in LEVELS
                                    else fallback["experience_level"])
    return {"name": str(data.get("name") or "")[:60], "skills": skills or fallback["skills"],
            "suggested_role": role, "experience_level": level,
            "summary": str(data.get("summary") or "")[:400], "method": "llm"}
