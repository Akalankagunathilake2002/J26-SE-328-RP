"""Read PDF, DOCX and TXT documents into plain text (TAF task 3)."""
import io
from pathlib import Path

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentError(ValueError):
    """The document cannot be read."""


def _read_pdf(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _read_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _read_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def extract_text(data: bytes, filename: str) -> str:
    """Return the raw text of an uploaded document; the extension picks the reader."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(f"Unsupported file type {extension!r}; use PDF, DOCX or TXT")
    try:
        if extension == ".pdf":
            return _read_pdf(data)
        if extension == ".docx":
            return _read_docx(data)
        return _read_txt(data)
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError(f"Could not read {filename}: {error}") from error


def read_file(path: Path) -> str:
    return extract_text(path.read_bytes(), path.name)
