import io

import pytest
from docx import Document

from app.evidence.rules import classify_evidence
from app.extraction.cleaner import clean_text
from app.extraction.matcher import SkillEntry, SkillMatcher, load_rules
from app.extraction.pipeline import extract_from_document, extract_from_text
from app.extraction.reader import DocumentError, extract_text
from app.extraction.sections import heading_type, split_sections
from app.schemas.evidence import EvidenceType, SourceType
from app.schemas.extraction import SectionType

CV = """JANE PERERA
Software Engineering Intern
Summary
Graduate interested in learning Kubernetes.
Skills
• Python, JS, React.js, C, Go
• MS Excel, communication
Experience
Built a REST API with
FastAPI and PostgreSQL.
Projects
Developed the dashboard in React.
Education
BSc in Information Technology, modules in Java.
Certifications
AWS Certified Cloud Practitioner
"""

ENTRIES = [
    SkillEntry("SK-1", "Python", ()),
    SkillEntry("SK-2", "JavaScript", ("JS",)),
    SkillEntry("SK-3", "React", ("React.js",)),
    SkillEntry("SK-4", "C", ()),
    SkillEntry("SK-5", "Go", ()),
    SkillEntry("SK-6", "Microsoft Excel", ("MS Excel", "Excel")),
    SkillEntry("SK-7", "communication principles", ("communication",)),
    SkillEntry("SK-8", "FastAPI", ()),
    SkillEntry("SK-9", "PostgreSQL", ()),
    SkillEntry("SK-10", "Java", ()),
    SkillEntry("SK-11", "Amazon Web Services", ("AWS",)),
    SkillEntry("SK-12", "Kubernetes", ()),
    SkillEntry("SK-13", "Ada", ()),
    SkillEntry("SK-14", "Science", ()),
]


@pytest.fixture
def matcher():
    return SkillMatcher(ENTRIES, {"ada": "exact_case", "science": "skip"})


def test_reads_txt_and_docx():
    assert extract_text("Python\n".encode("utf-8"), "cv.txt") == "Python\n"
    document = Document()
    document.add_paragraph("Skills")
    document.add_paragraph("Python")
    buffer = io.BytesIO()
    document.save(buffer)
    assert extract_text(buffer.getvalue(), "cv.DOCX") == "Skills\nPython"


def test_rejects_unsupported_and_broken_files():
    with pytest.raises(DocumentError):
        extract_text(b"x", "cv.png")
    with pytest.raises(DocumentError):
        extract_text(b"not a pdf", "cv.pdf")


def test_clean_text_fixes_bullets_encoding_duplicates_and_broken_lines():
    raw = "Company Name ï¼​ City\n• Built a REST API with\nFastAPI.\n\nFooter\nSkills\nFooter\nSkills\n"
    assert clean_text(raw) == (
        "Company Name City\nBuilt a REST API with FastAPI.\nFooter\nSkills\nSkills"
    )


def test_clean_text_never_joins_a_line_onto_a_heading():
    assert clean_text("Skills\npython, java") == "Skills\npython, java"


def test_heading_type_ignores_case_colons_and_ampersands():
    assert heading_type("WORK HISTORY:") == SectionType.EXPERIENCE
    assert heading_type("Education & Training") == SectionType.EDUCATION
    assert heading_type("Highlights") == SectionType.SKILLS
    assert heading_type("Built a REST API") is None


def test_split_sections_keeps_offsets_into_the_text():
    text = clean_text(CV)
    sections = split_sections(text)
    assert [s.type for s in sections] == [
        SectionType.HEADER, SectionType.SUMMARY, SectionType.SKILLS, SectionType.EXPERIENCE,
        SectionType.PROJECTS, SectionType.EDUCATION, SectionType.CERTIFICATIONS,
    ]
    for section in sections:
        assert text[section.start:section.start + len(section.text)] == section.text


def test_matcher_finds_names_and_aliases_ignoring_case(matcher):
    found = [(m.skill_id, m.matched_text) for m in matcher.find("python, JS and react.js")]
    assert found == [("SK-1", "python"), ("SK-2", "JS"), ("SK-3", "react.js")]


def test_short_terms_match_exact_case_whole_words_only(matcher):
    found = [m.matched_text for m in matcher.find("Go and C, go to grade c, Cisco, Golang")]
    assert found == ["Go", "C"]


def test_longest_match_wins(matcher):
    assert [m.skill_id for m in matcher.find("Microsoft Excel")] == ["SK-6"]
    assert [m.matched_text for m in matcher.find("Microsoft Excel")] == ["Microsoft Excel"]


def test_brackets_and_slashes_split_tokens(matcher):
    assert [m.skill_id for m in matcher.find("JavaScript(ES6)/Python")] == ["SK-2", "SK-1"]


def test_matcher_rules_exact_case_and_skip(matcher):
    assert [m.skill_id for m in matcher.find("Ada, ADA policy, Bachelor of Science")] == ["SK-13"]


def test_load_rules_rejects_unknown_rule(tmp_path):
    path = tmp_path / "matcher_rules.csv"
    path.write_text("term,rule,note\nLESS,lowercase,\n", encoding="utf-8")
    with pytest.raises(ValueError):
        load_rules(path)
    assert load_rules(tmp_path / "missing.csv") == {}


@pytest.mark.parametrize(
    ("source", "section", "sentence", "expected"),
    [
        (SourceType.CV, SectionType.SKILLS, "Python, Java", EvidenceType.CV_SELF_REPORTED),
        (SourceType.CV, SectionType.EXPERIENCE, "Built an API in FastAPI",
         EvidenceType.PROJECT_DEMONSTRATED),
        (SourceType.CV, SectionType.EXPERIENCE, "Python", EvidenceType.CV_SELF_REPORTED),
        (SourceType.CV, SectionType.EDUCATION, "Modules: Java", EvidenceType.ACADEMIC_EXPOSURE),
        (SourceType.CV, SectionType.CERTIFICATIONS, "AWS Certified",
         EvidenceType.CERTIFICATE_SUPPORTED),
        (SourceType.CV, SectionType.HEADER, "Software Engineer", EvidenceType.UNVERIFIED_MENTION),
        (SourceType.CV, SectionType.PROJECTS, "Interested in learning Kubernetes",
         EvidenceType.UNVERIFIED_MENTION),
        (SourceType.CERTIFICATE, SectionType.HEADER, "AWS", EvidenceType.CERTIFICATE_SUPPORTED),
        (SourceType.ACADEMIC_DOCUMENT, SectionType.OTHER, "Java",
         EvidenceType.ACADEMIC_EXPOSURE),
        (SourceType.PROJECT_DESCRIPTION, SectionType.HEADER, "React",
         EvidenceType.PROJECT_DEMONSTRATED),
    ],
)
def test_classify_evidence(source, section, sentence, expected):
    assert classify_evidence(source, section, sentence) == expected


def test_extract_from_text_end_to_end(matcher):
    result = extract_from_text(CV, matcher=matcher)
    by_skill = {s.skill_id: s for s in result.skills}

    assert by_skill["SK-12"].evidence_types == [EvidenceType.UNVERIFIED_MENTION]
    assert by_skill["SK-1"].evidence_types == [EvidenceType.CV_SELF_REPORTED]
    assert by_skill["SK-8"].evidence_types == [EvidenceType.PROJECT_DEMONSTRATED]
    assert by_skill["SK-3"].mention_count == 2
    assert by_skill["SK-3"].evidence_types == [
        EvidenceType.CV_SELF_REPORTED, EvidenceType.PROJECT_DEMONSTRATED,
    ]
    assert by_skill["SK-10"].evidence_types == [EvidenceType.ACADEMIC_EXPOSURE]
    assert by_skill["SK-11"].evidence_types == [EvidenceType.CERTIFICATE_SUPPORTED]

    fastapi = next(m for m in result.mentions if m.skill_id == "SK-8")
    assert fastapi.sentence == "Built a REST API with FastAPI and PostgreSQL."
    assert fastapi.method == "dictionary"


def test_mention_offsets_point_into_the_cleaned_text(matcher):
    result = extract_from_text(CV, matcher=matcher)
    text = clean_text(CV)
    for mention in result.mentions:
        assert text[mention.start:mention.end] == mention.matched_text


def test_extract_from_document_uses_the_real_dictionary():
    result = extract_from_document(b"Skills\nPython, JavaScript, SQL\n", "cv.txt")
    assert {"Python", "JavaScript"} <= {s.name for s in result.skills}


def test_skills_get_a_confidence_and_are_sorted_by_it(matcher):
    result = extract_from_text(CV, matcher=matcher)
    by_skill = {s.skill_id: s for s in result.skills}

    react = by_skill["SK-3"]  # listed and used in a project
    assert react.confidence_level == "weak"  # a CV alone is the student's own claim
    assert by_skill["SK-1"].confidence_level == "weak"  # only listed
    assert by_skill["SK-12"].confidence < by_skill["SK-1"].confidence < react.confidence
    assert [c.evidence_type for c in react.evidence] == [
        EvidenceType.PROJECT_DEMONSTRATED, EvidenceType.CV_SELF_REPORTED,
    ]
    confidences = [s.confidence for s in result.skills]
    assert confidences == sorted(confidences, reverse=True)
